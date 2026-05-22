# ingest_source_1.py
# Propósito: Script de ingesta para fuente 1 (_AIDA Ventures - Startups Benchmarks.xlsx)
# Ref: instrucciones_v2 — Sección 6

import pandas as pd
import numpy as np
from datetime import datetime
import openpyxl
from utils_ingest import (
    parse_range_value, parse_reference_date, detect_section_sector,
    standardize_stage, standardize_region, REGION_MAP, SECTOR_MAP, STAGE_MAP
)


OUTPUT_COLUMNS = [
    "record_id", "data_entry_date", "data_reference_date", "source_name",
    "source_type", "country", "region", "sector", "subsector", "round_stage",
    "metric_category", "metric_name", "metric_value_min", "metric_value_max",
    "metric_value_mid", "metric_unit", "currency", "investment_amount_usd",
    "valuation_pre_money_usd", "valuation_post_money_usd", "revenue_multiple",
    "growth_rate", "round_size_usd", "time_between_rounds_months",
    "graduation_rate", "deal_count", "startup_count", "notes",
    "confidence_level", "last_updated"
]

import os
# Detectar si estamos en raíz o en scripts/ingest
if os.path.exists("data/raw/sources"):
    SOURCE_FILE = "data/raw/sources/_AIDA Ventures - Startups Benchmarks.xlsx"
else:
    SOURCE_FILE = "../../data/raw/sources/_AIDA Ventures - Startups Benchmarks.xlsx"

SOURCE_NAME = "_AIDA_Ventures_-_Startups_Benchmarks.xlsx"


def create_record(base_dict, rownum, sheet_name):
    """Crea un registro completo con todas las columnas del schema."""
    record = {col: np.nan for col in OUTPUT_COLUMNS}

    record["record_id"] = f"S1_{sheet_name[:8].upper()}_{rownum:05d}"
    record["data_entry_date"] = datetime.now().date()
    record["source_name"] = SOURCE_NAME
    record["source_type"] = "public_report"
    record["metric_category"] = sheet_name
    record["confidence_level"] = "medium"
    record["last_updated"] = datetime.now().date()

    # Mergear base_dict
    for key, value in base_dict.items():
        if key in record and pd.notna(value):
            record[key] = value

    return record


def process_revenue_multiples(ws):
    """Procesa sheet 'Revenue Multiples'."""
    rows = []
    current_section = ("other", "")
    current_headers = None

    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        # Detectar sección
        if row[0] and isinstance(row[0], str) and "▸" in str(row[0]):
            current_section = detect_section_sector(str(row[0]), SECTOR_MAP)
            continue

        # Detectar headers
        if row[0] and "Stage" in str(row[0]):
            current_headers = row
            continue

        # Procesar datos
        if current_headers and row[0] and "Stage" not in str(row[0]):
            try:
                stage = standardize_stage(str(row[0]))
                global_multiple = str(row[1]) if len(row) > 1 and row[1] else ""
                source = str(row[2]) if len(row) > 2 and row[2] else ""
                date_text = str(row[3]) if len(row) > 3 and row[3] else ""

                if not global_multiple or global_multiple == "nan":
                    continue

                parsed = parse_range_value(global_multiple)

                record = create_record({
                    "country": "Global",
                    "region": "global",
                    "sector": current_section[0],
                    "subsector": current_section[1],
                    "round_stage": stage,
                    "metric_name": "revenue_multiple",
                    "metric_value_min": parsed["min"],
                    "metric_value_max": parsed["max"],
                    "metric_value_mid": parsed["mid"],
                    "metric_unit": parsed["unit"],
                    "currency": "USD" if parsed["unit"] == "usd" else "",
                    "revenue_multiple": parsed["mid"],
                    "data_reference_date": parse_reference_date(date_text),
                    "notes": parsed["notes"] if parsed["notes"] != global_multiple else "",
                }, row_idx, "Revenue_Multiples")

                rows.append(record)
            except Exception as e:
                pass

    return rows


def process_us_valuations(ws):
    """Procesa sheet 'US Valuations'."""
    rows = []
    current_section = ("other", "")
    current_headers = None

    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        if row[0] and isinstance(row[0], str) and "▸" in str(row[0]):
            current_section = detect_section_sector(str(row[0]), SECTOR_MAP)
            continue

        if row[0] and "Stage" in str(row[0]):
            current_headers = row
            continue

        if current_headers and row[0] and "Stage" not in str(row[0]):
            try:
                stage = standardize_stage(str(row[0]))
                median_val = str(row[1]) if len(row) > 1 and row[1] else ""

                if not median_val or median_val == "nan":
                    continue

                parsed = parse_range_value(median_val)

                record = create_record({
                    "country": "United States",
                    "region": "north_america",
                    "sector": current_section[0],
                    "subsector": current_section[1],
                    "round_stage": stage,
                    "metric_name": "valuation_pre_money_usd",
                    "metric_value_min": parsed["min"],
                    "metric_value_max": parsed["max"],
                    "metric_value_mid": parsed["mid"],
                    "metric_unit": "usd",
                    "currency": "USD",
                    "valuation_pre_money_usd": parsed["mid"],
                    "data_reference_date": pd.to_datetime("2025-01-01"),
                    "notes": parsed["notes"] if parsed["notes"] != median_val else "",
                }, row_idx, "US_Valuations")

                rows.append(record)
            except Exception as e:
                pass

    return rows


def process_latam_valuations(ws):
    """Procesa sheet 'LatAm Valuations'."""
    rows = []
    in_estimated_section = False

    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        if row[0] and isinstance(row[0], str) and "Estimated LATAM" in str(row[0]):
            in_estimated_section = True
            continue

        if not in_estimated_section:
            continue

        try:
            stage = str(row[0]).strip() if row[0] else ""
            valuation_str = str(row[1]) if len(row) > 1 and row[1] else ""

            if not stage or not valuation_str or "Stage" in stage or stage == "nan":
                continue

            stage = standardize_stage(stage)
            parsed = parse_range_value(valuation_str)

            record = create_record({
                "country": "LATAM",
                "region": "latam",
                "sector": "all_sectors",
                "round_stage": stage,
                "metric_name": "valuation_pre_money_usd",
                "metric_value_min": parsed["min"],
                "metric_value_max": parsed["max"],
                "metric_value_mid": parsed["mid"],
                "metric_unit": "usd",
                "currency": "USD",
                "valuation_pre_money_usd": parsed["mid"],
                "data_reference_date": pd.to_datetime("2024-12-31"),
                "notes": parsed["notes"] if parsed["notes"] != valuation_str else "",
            }, row_idx, "LatAm_Valuations")

            rows.append(record)
        except Exception as e:
            pass

    return rows


def process_time_between_rounds(ws):
    """Procesa sheet 'Time Between Rounds'."""
    rows = []

    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        try:
            transition = str(row[0]).strip() if row[0] else ""

            if "Transition" in transition or not transition or transition == "nan":
                continue

            if "LATAM-Specific" in transition:
                break

            # Extraer etapa origen del texto "Seed → Series A"
            if "→" in transition:
                source_stage = transition.split("→")[0].strip()
                source_stage = standardize_stage(source_stage)
            else:
                continue

            us_median = str(row[1]) if len(row) > 1 and row[1] else ""
            latam_median = str(row[2]) if len(row) > 2 and row[2] else ""

            # US row
            if us_median and us_median != "nan":
                parsed = parse_range_value(us_median)
                record = create_record({
                    "country": "United States",
                    "region": "north_america",
                    "round_stage": source_stage,
                    "metric_name": "time_between_rounds_months",
                    "metric_value_min": parsed["min"],
                    "metric_value_max": parsed["max"],
                    "metric_value_mid": parsed["mid"],
                    "metric_unit": "months",
                    "time_between_rounds_months": parsed["mid"],
                    "data_reference_date": pd.to_datetime("2024-12-31"),
                    "notes": parsed["notes"] if parsed["notes"] != us_median else "",
                }, row_idx, "Time_Between_Rounds")
                rows.append(record)

            # LATAM row
            if latam_median and latam_median != "nan":
                parsed = parse_range_value(latam_median)
                record = create_record({
                    "country": "LATAM",
                    "region": "latam",
                    "round_stage": source_stage,
                    "metric_name": "time_between_rounds_months",
                    "metric_value_min": parsed["min"],
                    "metric_value_max": parsed["max"],
                    "metric_value_mid": parsed["mid"],
                    "metric_unit": "months",
                    "time_between_rounds_months": parsed["mid"],
                    "data_reference_date": pd.to_datetime("2024-12-31"),
                    "notes": parsed["notes"] if parsed["notes"] != latam_median else "",
                }, row_idx, "Time_Between_Rounds")
                rows.append(record)
        except Exception as e:
            pass

    return rows


def process_graduation_rates(ws):
    """Procesa sheet 'Graduation Rates'."""
    rows = []
    in_section = False

    for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
        if row[0] and isinstance(row[0], str) and "Stage Graduation Rates" in str(row[0]):
            in_section = True
            continue

        if not in_section:
            continue

        try:
            stage = str(row[0]).strip() if row[0] else ""

            if not stage or "Graduation" in stage or stage == "nan" or "–" not in stage:
                continue

            us_rate = str(row[1]) if len(row) > 1 and row[1] else ""
            latam_rate = str(row[2]) if len(row) > 2 and row[2] else ""

            # Extraer etapa origen
            if "→" in stage:
                source_stage = stage.split("→")[0].strip()
                source_stage = standardize_stage(source_stage)
            else:
                continue

            # US row
            if us_rate and us_rate != "nan":
                parsed = parse_range_value(us_rate)
                record = create_record({
                    "country": "United States",
                    "region": "north_america",
                    "round_stage": source_stage,
                    "metric_name": "graduation_rate",
                    "metric_value_min": parsed["min"],
                    "metric_value_max": parsed["max"],
                    "metric_value_mid": parsed["mid"],
                    "metric_unit": "pct",
                    "graduation_rate": parsed["mid"],
                    "data_reference_date": pd.to_datetime("2024-12-31"),
                    "notes": parsed["notes"] if parsed["notes"] != us_rate else "",
                }, row_idx, "Graduation_Rates")
                rows.append(record)

            # LATAM row
            if latam_rate and latam_rate != "nan":
                parsed = parse_range_value(latam_rate)
                record = create_record({
                    "country": "LATAM",
                    "region": "latam",
                    "round_stage": source_stage,
                    "metric_name": "graduation_rate",
                    "metric_value_min": parsed["min"],
                    "metric_value_max": parsed["max"],
                    "metric_value_mid": parsed["mid"],
                    "metric_unit": "pct",
                    "graduation_rate": parsed["mid"],
                    "data_reference_date": pd.to_datetime("2024-12-31"),
                    "notes": parsed["notes"] if parsed["notes"] != latam_rate else "",
                }, row_idx, "Graduation_Rates")
                rows.append(record)
        except Exception as e:
            pass

    return rows


def run():
    """Función principal que procesa todos los sheets y retorna DataFrame."""
    try:
        wb = openpyxl.load_workbook(SOURCE_FILE)
    except Exception as e:
        print(f"ERROR: Cannot open {SOURCE_FILE}: {e}")
        return None

    all_rows = []

    # Revenue Multiples
    if "Revenue Multiples" in wb.sheetnames:
        ws = wb["Revenue Multiples"]
        all_rows.extend(process_revenue_multiples(ws))

    # US Valuations
    if "US Valuations" in wb.sheetnames:
        ws = wb["US Valuations"]
        all_rows.extend(process_us_valuations(ws))

    # LatAm Valuations
    if "LatAm Valuations" in wb.sheetnames:
        ws = wb["LatAm Valuations"]
        all_rows.extend(process_latam_valuations(ws))

    # Time Between Rounds
    if "Time Between Rounds" in wb.sheetnames:
        ws = wb["Time Between Rounds"]
        all_rows.extend(process_time_between_rounds(ws))

    # Graduation Rates
    if "Graduation Rates" in wb.sheetnames:
        ws = wb["Graduation Rates"]
        all_rows.extend(process_graduation_rates(ws))

    df = pd.DataFrame(all_rows)

    # Asegurar tipos correctos
    for col in ["metric_value_min", "metric_value_max", "metric_value_mid"]:
        df[col] = pd.to_numeric(df[col], errors="coerce")

    print(f"source_1: {len(df)} rows generadas")
    return df


if __name__ == "__main__":
    df = run()
    if df is not None:
        print(f"Total rows: {len(df)}")

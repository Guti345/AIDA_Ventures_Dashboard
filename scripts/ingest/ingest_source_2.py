# ingest_source_2.py
# Propósito: Script de ingesta para fuente 2 (_Metricas_Startups.xlsx)
# Ref: instrucciones_v2 — Sección 6

import pandas as pd
import numpy as np
from datetime import datetime
import openpyxl
import os

# Detectar si estamos en raíz o en scripts/ingest
if os.path.exists("data/raw/sources"):
    SOURCE_FILE = "data/raw/sources/_Metricas Startups.xlsx"
else:
    SOURCE_FILE = "../../data/raw/sources/_Metricas Startups.xlsx"

from utils_ingest import (
    parse_range_value, standardize_stage
)


OUTPUT_COLUMNS = [
    "record_id", "data_entry_date", "data_reference_date", "source_name",
    "source_type", "country", "region", "sector", "subsector", "round_stage",
    "metric_category", "metric_name", "metric_value_min", "metric_value_max",
    "metric_value_mid", "metric_unit", "currency", "investment_amount_usd",
    "valuation_pre_money_usd", "valuation_post_money_usd", "revenue_multiple",
    "growth_rate", "round_size_usd", "time_between_rounds_months",
    "graduation_rate", "deal_count", "startup_count", "notes",
    "confidence_level", "last_updated"
]

SOURCE_NAME = "_Metricas_Startups.xlsx"

METRIC_NAME_MAP = {
    "ARR (Annual Recurring Revenue)": "arr_usd",
    "MRR (Monthly Recurring Revenue)": "mrr_usd",
    "Crecimiento MoM (mes a mes)": "growth_mom",
    "Crecimiento YoY (año a año) en ARR": "growth_yoy",
    "Tiempo para alcanzar $1M ARR": "time_to_1m_arr_months",
    "Tiempo para alcanzar $10M ARR": "time_to_10m_arr_months",
    "Rule of 40 (crecimiento % + margen %)": "rule_of_40",
    "ARR por empleado": "arr_per_employee",
    "Burn Rate mensual (net)": "burn_rate_usd",
    "Runway objetivo": "runway_months",
    "Burn Multiple (burn neto / nuevo ARR)": "burn_multiple",
    "Eficiencia de capital (BEI)": "capital_efficiency_index",
    "Capital total levantado en ronda": "round_amount_usd",
    "Dilución típica por ronda": "dilution_pct",
    "Tiempo entre rondas (Seed → Serie A)": "time_between_rounds_months",
    "Churn Mensual (tasa de cancelación)": "churn_rate_monthly",
    "Churn Anual": "churn_rate_annual",
    "NRR (Net Revenue Retention)": "nrr",
    "GRR (Gross Revenue Retention)": "grr",
    "Retención de clientes": "customer_retention_rate",
    "NPS (Net Promoter Score)": "nps",
    "DAU/MAU ratio (para apps consumer)": "dau_mau_ratio",
    "CAC (Costo de Adquisición de Cliente)": "cac_usd",
    "LTV (Lifetime Value)": "ltv_usd",
    "LTV/CAC Ratio": "ltv_cac_ratio",
    "CAC Payback Period": "cac_payback_months",
    "Magic Number (eficiencia GTM)": "magic_number",
    "Gross Margin": "gross_margin",
    "Valuación pre-money mediana": "valuation_pre_money_usd",
    "Múltiplo ARR (para valuación)": "revenue_multiple",
    "TAM mínimo esperado": "tam_usd",
    "ARR para calificar a Serie A": "arr_threshold_series_a",
    "ARR para \"excepcionalmente fuerte\" Serie A": "arr_threshold_series_a_strong",
    "Valuación Seed (ARR >$250K)": "valuation_seed_arr_above_250k",
    "Valuación Seed (pre-revenue)": "valuation_seed_pre_revenue",
    "Tamaño del equipo (headcount)": "headcount",
    "MVP o producto lanzado": "product_stage",
    "Tasa de conversión Seed → Serie A": "conversion_rate_seed_to_series_a",
    "Número de clientes pagantes": "active_customers",
    "Duración del ciclo de ventas (B2B)": "sales_cycle_days",
    "ARR mínimo para \"seed estándar\" 2024": "arr_threshold_seed_standard",
}


def create_record(base_dict, rownum, sheet_name):
    """Crea un registro completo con todas las columnas del schema."""
    record = {col: np.nan for col in OUTPUT_COLUMNS}

    record["record_id"] = f"S2_{sheet_name[:8].upper()}_{rownum:05d}"
    record["data_entry_date"] = datetime.now().date()
    record["source_name"] = SOURCE_NAME
    record["source_type"] = "public_report"
    record["metric_category"] = sheet_name
    record["confidence_level"] = "medium"
    record["last_updated"] = datetime.now().date()

    # Mergear base_dict
    for key, value in base_dict.items():
        if key in record and pd.notna(value):
            record[key] = value

    return record


def process_sheet(ws, sheet_name):
    """Procesa un sheet con estructura: fila 1=headers, filas 2+=datos."""
    rows = []

    # Leer todas las filas
    all_rows_list = list(ws.iter_rows(values_only=True))

    if len(all_rows_list) < 3:
        return rows

    # Headers en fila 1 (índice 1)
    headers_row = all_rows_list[1]
    if not headers_row or len(headers_row) < 2:
        return rows

    # Encontrar índices de etapas
    stage_cols = {}
    for col_idx, header_val in enumerate(headers_row):
        if not header_val:
            continue
        header_str = str(header_val).strip().lower()
        if "pre-seed" in header_str or "preseed" in header_str:
            stage_cols["pre_seed"] = col_idx
        elif "seed" in header_str and "serie" not in header_str:
            stage_cols["seed"] = col_idx
        elif "serie a" in header_str or "series a" in header_str:
            stage_cols["series_a"] = col_idx

    # Procesar filas de datos (desde fila 2 en adelante)
    for row_idx, row in enumerate(all_rows_list[2:], start=2):
        metric_name = row[1] if len(row) > 1 and row[1] else None

        if not metric_name or str(metric_name).strip() == "":
            continue

        metric_name_str = str(metric_name).strip()

        # Mapear nombre de métrica
        mapped_name = METRIC_NAME_MAP.get(metric_name_str)
        if not mapped_name:
            # Fallback: convertir a snake_case
            mapped_name = metric_name_str.lower().replace(" ", "_").replace("(", "").replace(")", "")
            notes_prefix = "[unmapped] "
        else:
            notes_prefix = ""

        # Por cada etapa, crear un registro
        for stage, col_idx in stage_cols.items():
            if col_idx >= len(row):
                continue

            value = row[col_idx]
            if not value or str(value).strip() == "" or str(value).lower() == "n/a":
                continue

            value_str = str(value).strip()

            # Parsear el valor
            parsed = parse_range_value(value_str)

            record = create_record({
                "country": "United States",
                "region": "north_america",
                "sector": "all_sectors",
                "round_stage": stage,
                "metric_name": mapped_name,
                "metric_value_min": parsed["min"],
                "metric_value_max": parsed["max"],
                "metric_value_mid": parsed["mid"],
                "metric_unit": parsed["unit"],
                "currency": "USD" if parsed["unit"] == "usd" else "",
                "data_reference_date": pd.to_datetime("2025-01-01"),
                "notes": notes_prefix + parsed["notes"] if parsed["notes"] else notes_prefix,
            }, row_idx, sheet_name)

            rows.append(record)

    return rows


def run():
    """Función principal que procesa todos los sheets y retorna DataFrame."""
    try:
        wb = openpyxl.load_workbook(SOURCE_FILE)
    except Exception as e:
        print(f"ERROR: Cannot open {SOURCE_FILE}: {e}")
        return None

    all_rows = []

    # Procesar sheets especificados
    sheets_to_process = [
        "Income&Growth US",
        "Capital Efficiency US",
        "Retention&Customers",
        "Unit Economics",
        "Market&Valuation",
        "Operating Metrics",
    ]

    for sheet_name in sheets_to_process:
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            all_rows.extend(process_sheet(ws, sheet_name))

    df = pd.DataFrame(all_rows)

    # Asegurar tipos correctos
    for col in ["metric_value_min", "metric_value_max", "metric_value_mid"]:
        df[col] = pd.to_numeric(df[col], errors="coerce")

    print(f"source_2: {len(df)} rows generadas")
    return df


if __name__ == "__main__":
    df = run()
    if df is not None:
        print(f"Total rows: {len(df)}")

# ingest_source_3.py
# Propósito: Script de ingesta para fuente 3 (Fintech_Sectors.xlsx)
# Ref: instrucciones_v2 — Sección 6

import pandas as pd
import numpy as np
from datetime import datetime
import openpyxl
import os

# Detectar si estamos en raíz o en scripts/ingest
if os.path.exists("data/raw/sources"):
    SOURCE_FILE = "data/raw/sources/Fintech Sectors.xlsx"
else:
    SOURCE_FILE = "../../data/raw/sources/Fintech Sectors.xlsx"

from utils_ingest import (
    parse_range_value, standardize_region
)


OUTPUT_COLUMNS = [
    "record_id", "data_entry_date", "data_reference_date", "source_name",
    "source_type", "country", "region", "sector", "subsector", "round_stage",
    "metric_category", "metric_name", "metric_value_min", "metric_value_max",
    "metric_value_mid", "metric_unit", "currency", "investment_amount_usd",
    "valuation_pre_money_usd", "valuation_post_money_usd", "revenue_multiple",
    "growth_rate", "round_size_usd", "time_between_rounds_months",
    "graduation_rate", "deal_count", "startup_count", "notes",
    "confidence_level", "last_updated"
]

SOURCE_NAME = "Fintech_Sectors.xlsx"


def create_record(base_dict, rownum, sheet_name):
    """Crea un registro completo con todas las columnas del schema."""
    record = {col: np.nan for col in OUTPUT_COLUMNS}

    record["record_id"] = f"S3_{sheet_name[:8].upper()}_{rownum:05d}"
    record["data_entry_date"] = datetime.now().date()
    record["source_name"] = SOURCE_NAME
    record["source_type"] = "public_report"
    record["metric_category"] = sheet_name
    record["confidence_level"] = "medium"
    record["last_updated"] = datetime.now().date()

    # Mergear base_dict
    for key, value in base_dict.items():
        if key in record and pd.notna(value):
            record[key] = value

    return record


def process_investment_by_country(ws):
    """Procesa sheet 'Investment by Country 2024'."""
    rows = []
    all_rows_list = list(ws.iter_rows(values_only=True))

    # Headers en fila 0
    headers = all_rows_list[0] if all_rows_list else []

    # Procesar filas de datos
    for row_idx, row in enumerate(all_rows_list[1:], start=1):
        country = row[0] if len(row) > 0 and row[0] else None

        if not country or str(country).strip() == "":
            continue

        country_str = str(country).strip()
        investment = row[1] if len(row) > 1 and row[1] else ""
        deals = row[4] if len(row) > 4 and row[4] else ""
        num_fintechs = row[3] if len(row) > 3 and row[3] else ""
        growth_yoy = row[6] if len(row) > 6 and row[6] else ""

        region = standardize_region(country_str)

        # Investment row
        if investment and str(investment).strip():
            parsed = parse_range_value(str(investment))
            record = create_record({
                "country": country_str,
                "region": region,
                "sector": "fintech",
                "round_stage": "all_stages",
                "metric_name": "investment_amount_usd",
                "metric_value_min": parsed["min"],
                "metric_value_max": parsed["max"],
                "metric_value_mid": parsed["mid"],
                "metric_unit": "usd",
                "currency": "USD",
                "investment_amount_usd": parsed["mid"],
                "data_reference_date": pd.to_datetime("2024-12-31"),
                "notes": parsed["notes"] if parsed["notes"] != str(investment) else "",
            }, row_idx, "Investment_by_Country_2024")
            rows.append(record)

        # Deals row
        if deals and str(deals).strip():
            parsed = parse_range_value(str(deals))
            record = create_record({
                "country": country_str,
                "region": region,
                "sector": "fintech",
                "round_stage": "all_stages",
                "metric_name": "deal_count",
                "metric_value_min": parsed["min"],
                "metric_value_max": parsed["max"],
                "metric_value_mid": parsed["mid"],
                "metric_unit": "count",
                "deal_count": int(parsed["mid"]) if not np.isnan(parsed["mid"]) else np.nan,
                "data_reference_date": pd.to_datetime("2024-12-31"),
                "notes": parsed["notes"] if parsed["notes"] != str(deals) else "",
            }, row_idx, "Investment_by_Country_2024")
            rows.append(record)

        # Startups count row
        if num_fintechs and str(num_fintechs).strip():
            parsed = parse_range_value(str(num_fintechs))
            record = create_record({
                "country": country_str,
                "region": region,
                "sector": "fintech",
                "round_stage": "all_stages",
                "metric_name": "startup_count",
                "metric_value_min": parsed["min"],
                "metric_value_max": parsed["max"],
                "metric_value_mid": parsed["mid"],
                "metric_unit": "count",
                "startup_count": int(parsed["mid"]) if not np.isnan(parsed["mid"]) else np.nan,
                "data_reference_date": pd.to_datetime("2024-12-31"),
                "notes": parsed["notes"] if parsed["notes"] != str(num_fintechs) else "",
            }, row_idx, "Investment_by_Country_2024")
            rows.append(record)

        # Growth YoY row
        if growth_yoy and str(growth_yoy).strip():
            parsed = parse_range_value(str(growth_yoy))
            record = create_record({
                "country": country_str,
                "region": region,
                "sector": "fintech",
                "round_stage": "all_stages",
                "metric_name": "growth_yoy",
                "metric_value_min": parsed["min"],
                "metric_value_max": parsed["max"],
                "metric_value_mid": parsed["mid"],
                "metric_unit": "pct",
                "growth_rate": parsed["mid"],
                "data_reference_date": pd.to_datetime("2024-12-31"),
                "notes": parsed["notes"] if parsed["notes"] != str(growth_yoy) else "",
            }, row_idx, "Investment_by_Country_2024")
            rows.append(record)

    return rows


def process_latam_subsectors(ws):
    """Procesa sheet 'LATAM Subsectors 2024'."""
    rows = []
    all_rows_list = list(ws.iter_rows(values_only=True))

    for row_idx, row in enumerate(all_rows_list[1:], start=1):
        subsector = row[0] if len(row) > 0 and row[0] else None

        if not subsector or str(subsector).strip() == "":
            continue

        subsector_str = str(subsector).strip().lower().replace(" ", "_").replace("&", "and")
        investment = row[1] if len(row) > 1 and row[1] else ""
        num_startups = row[3] if len(row) > 3 and row[3] else ""

        # Investment row
        if investment and str(investment).strip():
            parsed = parse_range_value(str(investment))
            record = create_record({
                "country": "LATAM",
                "region": "latam",
                "sector": "fintech",
                "subsector": subsector_str,
                "metric_name": "investment_amount_usd",
                "metric_value_min": parsed["min"],
                "metric_value_max": parsed["max"],
                "metric_value_mid": parsed["mid"],
                "metric_unit": "usd",
                "currency": "USD",
                "investment_amount_usd": parsed["mid"],
                "data_reference_date": pd.to_datetime("2024-12-31"),
                "notes": parsed["notes"] if parsed["notes"] != str(investment) else "",
            }, row_idx, "LATAM_Subsectors_2024")
            rows.append(record)

        # Startups count row
        if num_startups and str(num_startups).strip():
            parsed = parse_range_value(str(num_startups))
            record = create_record({
                "country": "LATAM",
                "region": "latam",
                "sector": "fintech",
                "subsector": subsector_str,
                "metric_name": "startup_count",
                "metric_value_min": parsed["min"],
                "metric_value_max": parsed["max"],
                "metric_value_mid": parsed["mid"],
                "metric_unit": "count",
                "startup_count": int(parsed["mid"]) if not np.isnan(parsed["mid"]) else np.nan,
                "data_reference_date": pd.to_datetime("2024-12-31"),
                "notes": parsed["notes"] if parsed["notes"] != str(num_startups) else "",
            }, row_idx, "LATAM_Subsectors_2024")
            rows.append(record)

    return rows


def process_usa_subsectors(ws):
    """Procesa sheet 'USA Subsectors 2024'."""
    rows = []
    all_rows_list = list(ws.iter_rows(values_only=True))

    for row_idx, row in enumerate(all_rows_list[1:], start=1):
        subsector = row[0] if len(row) > 0 and row[0] else None

        if not subsector or str(subsector).strip() == "":
            continue

        subsector_str = str(subsector).strip().lower().replace(" ", "_").replace("&", "and")
        investment = row[1] if len(row) > 1 and row[1] else ""
        num_startups = row[3] if len(row) > 3 and row[3] else ""

        # Investment row
        if investment and str(investment).strip():
            parsed = parse_range_value(str(investment))
            record = create_record({
                "country": "United States",
                "region": "north_america",
                "sector": "fintech",
                "subsector": subsector_str,
                "metric_name": "investment_amount_usd",
                "metric_value_min": parsed["min"],
                "metric_value_max": parsed["max"],
                "metric_value_mid": parsed["mid"],
                "metric_unit": "usd",
                "currency": "USD",
                "investment_amount_usd": parsed["mid"],
                "data_reference_date": pd.to_datetime("2024-12-31"),
                "notes": parsed["notes"] if parsed["notes"] != str(investment) else "",
            }, row_idx, "USA_Subsectors_2024")
            rows.append(record)

        # Startups count row
        if num_startups and str(num_startups).strip():
            parsed = parse_range_value(str(num_startups))
            record = create_record({
                "country": "United States",
                "region": "north_america",
                "sector": "fintech",
                "subsector": subsector_str,
                "metric_name": "startup_count",
                "metric_value_min": parsed["min"],
                "metric_value_max": parsed["max"],
                "metric_value_mid": parsed["mid"],
                "metric_unit": "count",
                "startup_count": int(parsed["mid"]) if not np.isnan(parsed["mid"]) else np.nan,
                "data_reference_date": pd.to_datetime("2024-12-31"),
                "notes": parsed["notes"] if parsed["notes"] != str(num_startups) else "",
            }, row_idx, "USA_Subsectors_2024")
            rows.append(record)

    return rows


def run():
    """Función principal que procesa todos los sheets y retorna DataFrame."""
    try:
        wb = openpyxl.load_workbook(SOURCE_FILE)
    except Exception as e:
        print(f"ERROR: Cannot open {SOURCE_FILE}: {e}")
        return None

    all_rows = []

    # Investment by Country
    if "Investment by Country 2024" in wb.sheetnames:
        ws = wb["Investment by Country 2024"]
        all_rows.extend(process_investment_by_country(ws))

    # LATAM Subsectors
    if "LATAM Subsectors 2024" in wb.sheetnames:
        ws = wb["LATAM Subsectors 2024"]
        all_rows.extend(process_latam_subsectors(ws))

    # USA Subsectors
    if "USA Subsectors 2024" in wb.sheetnames:
        ws = wb["USA Subsectors 2024"]
        all_rows.extend(process_usa_subsectors(ws))

    df = pd.DataFrame(all_rows)

    # Asegurar tipos correctos
    for col in ["metric_value_min", "metric_value_max", "metric_value_mid"]:
        df[col] = pd.to_numeric(df[col], errors="coerce")

    print(f"source_3: {len(df)} rows generadas")
    return df


if __name__ == "__main__":
    df = run()
    if df is not None:
        print(f"Total rows: {len(df)}")
